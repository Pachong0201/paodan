# -*- coding: utf-8 -*-
"""V3.1 重要爆料邮件 Excel 台账测试（>=20项，全部使用临时目录）。"""
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.models import AttachmentDoc, EmailDocument, FinalScore, LLMResult, PatternHit, RuleResult, ScreeningRecord
from app.reports.excel_register import (
    HEADERS,
    MANUAL_DEFAULTS,
    MANUAL_FIELDS,
    ImportantEmailRegister,
)


def make_record(
    priority="A",
    score=82.0,
    message_id="mid-test-001@example.com",
    body_hash="hash-001",
    subject="爆料：某议员办公室疑协助建商处理建照",
    sender="whistle@example.com",
    date="Wed, 02 Sep 2026 10:00:00 +0800",
    categories=("A03", "A04"),
    subcats=(),
    persons=("林某某", "王某某"),
    orgs=("某市府",),
    companies=("某建商",),
    projects=("某建案",),
    one_line="邮件指称议员办公室协助建商处理建照并收受顾问费。",
    reason="含汇款与LINE纪录，值得记者核验。",
    patterns=("P04", "P05"),
    evidence=("汇款单", "LINE对话"),
    money_raw=("80万元",),
    new_info=("新增80万元汇款纪录",),
    old_info=(),
    vts=("核验80万元汇款真实性", "查询公司实际控制人", "对照建照审批时间"),
    primary="R5",
    secondary=("R3",),
    risks=("EVIDENCE_AUTHENTICITY",),
    formal=True,
    formal_types=("PROSECUTOR",),
    media=(("镜周刊", 97), ("TVBS", 91)),
    actors=(("黃國昌", 88),),
    amps=(("黃揚明", 80),),
    formals=(("台北地檢署", 90),),
    attachments=(("汇款截图.png", "LINE对话.pdf", "建照时间表.xlsx") if True else ()),
    source_path="data/inbox/sample_case2.eml",
    email_id="",
):
    eid = email_id or (message_id or body_hash)
    atts = [AttachmentDoc(filename=n, file_type="pdf") for n in (attachments or [])]
    email = EmailDocument(
        email_id=eid, message_id=message_id or "", subject=subject, sender=sender,
        date=date, body_text="正文 " + subject, body_hash=body_hash or "",
        attachments=atts, source_path=source_path,
    )
    pats = [PatternHit(pattern_id=p, category="A03", name=f"模式{p}") for p in (patterns or [])]
    money = [{"raw": m, "amount": 800000, "currency": "TWD"} for m in (money_raw or [])]
    rule = RuleResult(
        matched_categories=list(categories or []),
        matched_patterns=pats,
        money=money,
        entities=[{"text": persons[0] if persons else "x", "type": "PERSON", "count": 1}] if persons else [],
        target_persons_found=list(persons or []),
        target_orgs_found=list(orgs or []),
    )
    llm = LLMResult(
        categories=list(categories or []),
        subcategories=list(subcats or []),
        target_persons=list(persons or []),
        target_organizations=list(orgs or []),
        related_entities={"companies": list(companies or [])},
        projects_or_cases=list(projects or []),
        money_or_benefits=list(money_raw or []),
        evidence_items=list(evidence or []),
        new_information=list(new_info or []),
        known_old_information=list(old_info or []),
        verification_targets=list(vts or []),
        one_sentence_summary=one_line,
        reason_for_attention=reason,
        llm_status="template",
    )
    fs = FinalScore(final_score=float(score), priority=priority)
    rel = {
        "primary_route": primary or "",
        "primary_route_name": "",
        "secondary_routes": list(secondary or []),
        "avoid_routes": [],
        "route_confidence": 0.8,
        "prepublication_verification_required": True,
        "verification_before_release": [],
        "formal_referral_recommended": bool(formal),
        "formal_referral_type": list(formal_types or []),
        "release_risks": list(risks or []),
        "reason": "测试理由",
        "recommended_release_sequence": [],
    }
    named = {
        "recommended_media": [{"entity_id": f"m{i}", "name": n, "fit_score": s} for i, (n, s) in enumerate(media or [])],
        "recommended_disclosure_actors": [{"entity_id": f"a{i}", "name": n, "fit_score": s} for i, (n, s) in enumerate(actors or [])],
        "recommended_amplifiers": [{"entity_id": f"p{i}", "name": n, "fit_score": s} for i, (n, s) in enumerate(amps or [])],
        "recommended_formal_channels": [{"entity_id": f"f{i}", "name": n, "fit_score": s} for i, (n, s) in enumerate(formals or [])],
        "recommended_platforms": [],
        "local_channels": [],
    }
    rec = ScreeningRecord(email_id=eid, email=email, rule=rule, llm=llm, score=fs,
                          summary_zh=one_line, verification_targets=list(vts or []),
                          release_recommendation=rel, named_channel_recommendation=named)
    return rec


def _new_register(tmp_path, **kw):
    params = dict(path=tmp_path / "important_email_register.xlsx", enabled=True,
                  min_score=60.0, priorities=["S", "A", "B"], backup_before_batch=False)
    params.update(kw)
    return ImportantEmailRegister(**params)


def _read_rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    ws = wb["重要爆料邮件"]
    header = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        d = {header[c]: ws.cell(row=r, column=c + 1).value for c in range(len(header))}
        rows.append(d)
    return wb, ws, header, rows


# ---------- Test 01-05 门禁 ----------
def test_01_s_register(tmp_path):
    reg = _new_register(tmp_path)
    stats = reg.process_batch([make_record(priority="S", score=96)])
    assert stats["added"] == 1
    _, _, _, rows = _read_rows(reg.path)
    assert len(rows) == 1 and rows[0]["优先级"] == "S"


def test_02_a_register(tmp_path):
    reg = _new_register(tmp_path)
    stats = reg.process_batch([make_record(priority="A", score=82)])
    assert stats["added"] == 1


def test_03_b_register_default(tmp_path):
    reg = _new_register(tmp_path)
    stats = reg.process_batch([make_record(priority="B", score=65)])
    assert stats["added"] == 1


def test_04_c_not_register_default(tmp_path):
    reg = _new_register(tmp_path)
    stats = reg.process_batch([make_record(priority="C", score=50)])
    assert stats["added"] == 0
    assert not reg.path.exists() or _read_rows(reg.path)[3] == []


def test_05_d_not_register(tmp_path):
    reg = _new_register(tmp_path)
    stats = reg.process_batch([make_record(priority="D", score=20)])
    assert stats["added"] == 0


def test_06_score_threshold_config(tmp_path):
    # C 70分：min_score=60 应进（OR 逻辑），min_score=75 不应进
    reg60 = _new_register(tmp_path / "a", min_score=60.0)
    reg60.path.parent.mkdir(parents=True, exist_ok=True)
    # 用不同目录隔离
    reg60 = ImportantEmailRegister(path=tmp_path / "a.xlsx", min_score=60.0, priorities=["S", "A", "B"], backup_before_batch=False)
    assert reg60.should_register(make_record(priority="C", score=70)) is True
    reg75 = ImportantEmailRegister(path=tmp_path / "b.xlsx", min_score=75.0, priorities=["S", "A", "B"], backup_before_batch=False)
    assert reg75.should_register(make_record(priority="C", score=70)) is False
    assert reg75.should_register(make_record(priority="C", score=80)) is True
    # S/A/B 优先级通道不受高阈值影响
    assert reg75.should_register(make_record(priority="B", score=65)) is True


def test_07_message_id_dedup(tmp_path):
    reg = _new_register(tmp_path)
    rec = make_record(priority="A", score=82, message_id="dup-mid@example.com", body_hash="h1")
    assert reg.process_batch([rec])["added"] == 1
    assert reg.process_batch([rec])["added"] == 0
    _, _, _, rows = _read_rows(reg.path)
    assert len(rows) == 1


def test_08_no_message_id_hash_dedup(tmp_path):
    reg = _new_register(tmp_path)
    r1 = make_record(priority="A", score=82, message_id="", body_hash="same-hash-xyz")
    r2 = make_record(priority="A", score=82, message_id="", body_hash="same-hash-xyz")
    assert reg.process_batch([r1])["added"] == 1
    assert reg.process_batch([r2])["added"] == 0
    _, _, _, rows = _read_rows(reg.path)
    assert len(rows) == 1
    # 不同 hash 应新增
    r3 = make_record(priority="A", score=82, message_id="", body_hash="other-hash")
    assert reg.process_batch([r3])["added"] == 1
    _, _, _, rows = _read_rows(reg.path)
    assert len(rows) == 2


def test_09_score_change_updates_no_new_row(tmp_path):
    reg = _new_register(tmp_path)
    r1 = make_record(priority="A", score=82, message_id="chg@example.com", body_hash="h")
    assert reg.process_batch([r1])["added"] == 1
    r2 = make_record(priority="S", score=94, message_id="chg@example.com", body_hash="h")
    stats = reg.process_batch([r2])
    assert stats["added"] == 0 and stats["updated"] == 1
    _, _, _, rows = _read_rows(reg.path)
    assert len(rows) == 1
    assert rows[0]["优先级"] == "S"
    assert float(rows[0]["最终评分"]) == 94


def test_10_manual_note_protected(tmp_path):
    from openpyxl import load_workbook
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=82, message_id="m10@example.com")])
    # 模拟记者手工修改
    wb = load_workbook(str(reg.path))
    ws = wb["重要爆料邮件"]
    header = [c.value for c in ws[1]]
    note_col = header.index("记者备注") + 1
    ws.cell(row=2, column=note_col).value = "已联系来源"
    wb.save(str(reg.path))
    # 重跑
    reg2 = _new_register(tmp_path)
    reg2.path = reg.path
    stats = reg2.process_batch([make_record(priority="S", score=94, message_id="m10@example.com")])
    assert stats["updated"] == 1
    _, _, _, rows = _read_rows(reg.path)
    assert rows[0]["记者备注"] == "已联系来源"
    assert rows[0]["优先级"] == "S"  # 系统字段已更新


def test_11_status_protected(tmp_path):
    from openpyxl import load_workbook
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=82, message_id="m11@example.com")])
    wb = load_workbook(str(reg.path))
    ws = wb["重要爆料邮件"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("处理状态") + 1).value = "跟进中"
    ws.cell(row=2, column=header.index("责任编辑") + 1).value = "张编辑"
    ws.cell(row=2, column=header.index("人工标签") + 1).value = "高价值"
    wb.save(str(reg.path))
    reg2 = _new_register(tmp_path)
    reg2.path = reg.path
    reg2.process_batch([make_record(priority="S", score=95, message_id="m11@example.com")])
    _, _, _, rows = _read_rows(reg.path)
    assert rows[0]["处理状态"] == "跟进中"
    assert rows[0]["责任编辑"] == "张编辑"
    assert rows[0]["人工标签"] == "高价值"
    assert rows[0]["优先级"] == "S"


def test_12_v2_fields(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=85, primary="R5", secondary=("R3", "R2"))])
    _, _, _, rows = _read_rows(reg.path)
    assert "R5" not in str(rows[0]["V2推荐首发类型"])
    assert "正式检举优先型" in str(rows[0]["V2推荐首发类型"])
    assert "R3" not in str(rows[0]["V2备选渠道"]) and "R2" not in str(rows[0]["V2备选渠道"])
    assert "深度调查报道型" in str(rows[0]["V2备选渠道"])
    assert "媒体独家型" in str(rows[0]["V2备选渠道"])
    assert rows[0]["风险提示"] != ""
    # 正式检举类型写中文名
    assert "PROSECUTOR" not in str(rows[0]["风险提示"])
    assert "检察机关" in str(rows[0]["风险提示"])


def test_13_v3_fields(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=85)])
    _, _, _, rows = _read_rows(reg.path)
    assert "镜周刊" in str(rows[0]["V3推荐媒体"])
    assert "97" in str(rows[0]["V3推荐媒体"])
    assert "黃國昌" in str(rows[0]["V3推荐揭弊人物"])
    assert "黃揚明" in str(rows[0]["V3推荐放大者"])
    assert "台北地檢署" in str(rows[0]["V3正式渠道"])


def test_14_attachments_multi(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=85,
                                   attachments=("A.png", "B.pdf", "C.xlsx"))])
    _, _, _, rows = _read_rows(reg.path)
    assert rows[0]["附件数量"] == 3
    assert "A.png" in str(rows[0]["附件名称"]) and "；" in str(rows[0]["附件名称"])


def test_15_verification_multiline(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=85,
                                   vts=("核验80万元汇款真实性", "查询公司实际控制人"))])
    _, ws, _, rows = _read_rows(reg.path)
    txt = str(rows[0]["核查重点"])
    assert "1." in txt and "2." in txt and "\n" in txt
    # wrap_text
    header = [c.value for c in ws[1]]
    col = header.index("核查重点") + 1
    assert ws.cell(row=2, column=col).alignment.wrap_text is True


def test_16_auto_create(tmp_path):
    p = tmp_path / "sub" / "important_email_register.xlsx"
    reg = ImportantEmailRegister(path=p, backup_before_batch=False)
    assert not p.exists()
    reg.process_batch([make_record(priority="S", score=96)])
    assert p.exists()
    from openpyxl import load_workbook
    wb = load_workbook(str(p))
    assert "重要爆料邮件" in wb.sheetnames


def test_17_append_existing(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=82, message_id="a@x")])
    reg2 = _new_register(tmp_path)
    reg2.path = reg.path
    reg2.process_batch([make_record(priority="A", score=83, message_id="b@x")])
    _, _, _, rows = _read_rows(reg.path)
    assert len(rows) == 2
    assert [r["编号"] for r in rows] == [1, 2]


def test_18_occupied_no_crash(tmp_path, monkeypatch):
    reg = _new_register(tmp_path)
    rec = make_record(priority="A", score=82)
    # 模拟保存时被占用
    import openpyxl.workbook.workbook as _wbmod
    orig_save = _wbmod.Workbook.save

    def _boom(self, *a, **k):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr(_wbmod.Workbook, "save", _boom)
    stats = reg.process_batch([rec])  # 不得抛异常
    assert stats["added"] == 1 or stats["updated"] == 1
    # pending 应生成
    pending = reg.path.parent / "pending_excel_register.jsonl"
    assert pending.exists()
    monkeypatch.setattr(_wbmod.Workbook, "save", orig_save)


def test_19_corrupt_not_silent(tmp_path):
    p = tmp_path / "important_email_register.xlsx"
    p.write_bytes(b"not a real xlsx \x00\x01\x02")
    reg = ImportantEmailRegister(path=p, backup_before_batch=False)
    reg.process_batch([make_record(priority="A", score=82)])
    corrupts = list(tmp_path.glob("important_email_register.corrupt.*.xlsx"))
    assert len(corrupts) >= 1
    assert p.exists()
    _, _, _, rows = _read_rows(p)
    assert len(rows) == 1


def test_20_batch_100_single_save(tmp_path):
    reg = _new_register(tmp_path)
    recs = [make_record(priority="A" if i % 2 == 0 else "B", score=80,
                         message_id=f"batch-{i:03d}@x", body_hash=f"h{i}") for i in range(100)]
    stats = reg.process_batch(recs)
    assert stats["added"] == 100
    assert reg._save_count == 1
    _, _, _, rows = _read_rows(reg.path)
    assert len(rows) == 100


def test_21_format_gates(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="S", score=96)])
    from openpyxl import load_workbook
    wb = load_workbook(str(reg.path))
    ws = wb["重要爆料邮件"]
    header = [c.value for c in ws[1]]
    assert header == HEADERS
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws["A1"].font.bold is True
    # 重点列宽
    from openpyxl.utils import get_column_letter
    subj_col = HEADERS.index("邮件主题") + 1
    assert ws.column_dimensions[get_column_letter(subj_col)].width >= 30
    # 长文本 wrap
    seventy = HEADERS.index("一句话摘要") + 1
    assert ws.cell(row=2, column=seventy).alignment.wrap_text is True
    # 编号递增 + 人工字段存在
    assert ws.cell(row=2, column=1).value == 1
    for mf in MANUAL_FIELDS:
        assert mf in header
    # 默认处理状态
    assert ws.cell(row=2, column=header.index("处理状态") + 1).value == "待看"


def test_22_disabled_no_file(tmp_path):
    p = tmp_path / "important_email_register.xlsx"
    reg = ImportantEmailRegister(path=p, enabled=False)
    stats = reg.process_batch([make_record(priority="S", score=96)])
    assert not p.exists()
    assert stats["added"] == 0


def test_23_custom_priorities(tmp_path):
    reg = ImportantEmailRegister(path=tmp_path / "x.xlsx", priorities=["S"], min_score=100, backup_before_batch=False)
    assert reg.should_register(make_record(priority="S", score=95)) is True
    assert reg.should_register(make_record(priority="A", score=96)) is False
    assert reg.should_register(make_record(priority="B", score=99)) is False


def test_24_joiner_and_no_json_dump(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=85, categories=("A03", "A04"),
                                   persons=("林某某", "王某某"), patterns=("P04", "P05"))])
    _, _, _, rows = _read_rows(reg.path)
    assert "；" in str(rows[0]["主要类别"])
    assert "；" in str(rows[0]["涉及人物"])
    assert "；" in str(rows[0]["命中Pattern"])
    for v in rows[0].values():
        s = str(v or "")
        assert not (s.startswith("{") and s.endswith("}")), f"不应直接 dump JSON: {s[:80]}"


def test_28_no_type_codes_in_register(tmp_path):
    """台账不出现 A/P/R 类型代码，一律自然语言（未知代码除外）。"""
    import re
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=85, categories=("A03", "A04"),
                                   persons=("林某某", "王某某"), patterns=("P04", "P05"),
                                   one_line="邮件指称林某某涉A03相关事项，P04命中。",
                                   reason="A04值得关注，建议R5处理。",
                                   primary="R5", secondary=("R3",),
                                   formal_types=("PROSECUTOR", "CONTROL_YUAN"))])
    _, _, _, rows = _read_rows(reg.path)
    r = rows[0]
    assert "收贿、索贿及职务对价" in str(r["主要类别"])
    assert "土地开发" in str(r["主要类别"])
    assert "A03" not in str(r["主要类别"]) and "A04" not in str(r["主要类别"])
    assert "厂商金钱职务对价" in str(r["命中Pattern"])
    assert "P04" not in str(r["命中Pattern"]) and "P05" not in str(r["命中Pattern"])
    assert "A03" not in str(r["一句话摘要"]) and "P04" not in str(r["一句话摘要"])
    assert "正式检举优先型" in str(r["为什么值得看"]) and "R5" not in str(r["为什么值得看"])
    code_re = re.compile(r"(?<![A-Za-z0-9])[APR](0[1-9]|1[0-9]|20|[1-6])(?![0-9])")
    for col in ("主要类别", "子类别", "命中Pattern", "一句话摘要", "为什么值得看",
                "V2推荐首发类型", "V2备选渠道", "风险提示", "核查重点"):
        assert not code_re.search(str(r[col] or "")), f"{col}: {r[col]}"


def test_25_numbering_stable_on_update(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=82, message_id="n1@x")])
    reg2 = _new_register(tmp_path); reg2.path = reg.path
    reg2.process_batch([make_record(priority="A", score=83, message_id="n2@x")])
    reg3 = _new_register(tmp_path); reg3.path = reg.path
    reg3.process_batch([make_record(priority="S", score=95, message_id="n1@x")])
    _, _, _, rows = _read_rows(reg.path)
    by_mid = {r["Message-ID"]: r for r in rows}
    assert by_mid["n1@x"]["编号"] == 1
    assert by_mid["n2@x"]["编号"] == 2


def test_26_source_path_and_hash_columns(tmp_path):
    reg = _new_register(tmp_path)
    reg.process_batch([make_record(priority="A", score=85, message_id="src@x", body_hash="abc123",
                                   source_path="data/inbox/sample_case2.eml")])
    _, _, _, rows = _read_rows(reg.path)
    assert rows[0]["Message-ID"] == "src@x"
    assert rows[0]["Content Hash"] == "abc123"
    assert "sample_case2.eml" in str(rows[0]["原始邮件路径"])
    assert "Base64" not in str(rows[0]["附件名称"])


def test_27_register_key_logic(tmp_path):
    reg = _new_register(tmp_path)
    r_mid = make_record(priority="A", score=82, message_id="k@x", body_hash="h1")
    r_no = make_record(priority="A", score=82, message_id="", body_hash="h1")
    assert reg.build_register_key(r_mid).startswith("MID:")
    assert reg.build_register_key(r_no).startswith("HASH:")
    # 不得仅依赖 subject/sender/date：同主题不同 hash 应不同 key
    r_a = make_record(priority="A", score=82, message_id="", body_hash="ha", subject="同主题")
    r_b = make_record(priority="A", score=82, message_id="", body_hash="hb", subject="同主题")
    assert reg.build_register_key(r_a) != reg.build_register_key(r_b)
